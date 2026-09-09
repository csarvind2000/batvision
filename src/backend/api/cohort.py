"""
Cohort view over analysed BAT cases.

Population work means reading across subjects, not one at a time: this flattens
every case that has produced results into a single row of volumes, so the set
can be read as one sheet, sorted, and exported for statistics.

The numbers are taken from each case's ``bat_metrics.json`` — the file the AI
service writes — rather than recomputed here, so the sheet and the review screen
can never disagree.

A caveat that matters when reading these columns across subjects: the 3- and
4-class volumes are *percentile* bands of the fat-fraction distribution taken
inside each subject's own mask (cut at the 20th, 60th and 80th centiles). Each
class therefore holds a fixed share of that subject's mask by construction —
20/40/40 and 20/40/20/20 — so the per-class volumes are all proportional to the
total and carry no between-subject information the total does not already
carry. Comparing class volumes across a cohort compares mask sizes. Absolute
fat-fraction thresholds would be needed for the class split itself to vary
between subjects.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import os
from datetime import datetime
from statistics import mean, pstdev
from typing import Iterable, Optional

logger = logging.getLogger(__name__)

# Identity of the row. Kept deliberately short: a cohort sheet is read for its
# measurements, and every extra identifier column pushes them off screen.
BASE_COLUMNS = (
    ("case_id", "Subject"),
    ("patient_id", "Patient ID"),
    ("patient_name", "Patient name"),
    ("status", "Status"),
    ("created_at", "Added"),
)

# The measurements. Class names match the review panel and bat_metrics.json.
VOLUME_COLUMNS = (
    ("binary_total_ml", "BAT total (mL)"),
    ("c3_muscle_ml", "3-class muscle (mL)"),
    ("c3_brownfat_ml", "3-class brown fat (mL)"),
    ("c3_mixwhite_ml", "3-class mix+white (mL)"),
    ("c3_total_ml", "3-class total (mL)"),
    ("c4_muscle_ml", "4-class muscle (mL)"),
    ("c4_brownfat_ml", "4-class brown fat (mL)"),
    ("c4_mixfat_ml", "4-class mixed fat (mL)"),
    ("c4_whitefat_ml", "4-class white fat (mL)"),
    ("c4_total_ml", "4-class total (mL)"),
)

# Acquisition/QC facts that decide whether two rows are comparable at all.
QC_COLUMNS = (
    ("voxel_volume_ml", "Voxel volume (mL)"),
    ("pred_voxels", "BAT voxels"),
)

ALL_COLUMNS = BASE_COLUMNS + VOLUME_COLUMNS + QC_COLUMNS

# Columns the summary block reports n / mean / SD / min / max over.
SUMMARY_KEYS = tuple(key for key, _ in VOLUME_COLUMNS) + ("voxel_volume_ml",)


def _num(value) -> Optional[float]:
    """A metrics value as a float, or None when absent or not numeric."""
    if value is None or value == "":
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return out if out == out else None  # drop NaN


def _round(value: Optional[float], places: int = 3) -> Optional[float]:
    return None if value is None else round(value, places)


def read_metrics(out_dir: Optional[str]) -> Optional[dict]:
    """``bat_metrics.json`` from a case's output directory, or None."""
    if not out_dir:
        return None
    path = os.path.join(out_dir, "bat_metrics.json")
    if not os.path.isfile(path):
        return None
    try:
        with open(path, "r") as handle:
            return json.load(handle)
    except (OSError, ValueError) as error:
        logger.warning("Unreadable metrics at %s: %s", path, error)
        return None


def case_to_row(case, out_dir: Optional[str], metrics: Optional[dict]) -> dict:
    """One flat record for a case. Measurement cells are None when unavailable."""
    volumes = (metrics or {}).get("volumes") or {}
    stats = (metrics or {}).get("stats") or {}
    c3 = volumes.get("class3_breakdown_ml") or {}
    c4 = volumes.get("class4_breakdown_ml") or {}

    created = getattr(case, "created_at", None)
    row = {
        "id": case.pk,
        "case_id": getattr(case, "case_id", "") or "",
        "patient_id": getattr(case, "patient_id", None) or "",
        "patient_name": getattr(case, "patient_name", None) or "",
        "status": getattr(case, "status", "") or "",
        "created_at": created.isoformat() if isinstance(created, datetime) else "",
        "out_dir": out_dir or "",
        "has_metrics": bool(metrics),

        "binary_total_ml": _round(_num(volumes.get("binary_total_ml"))),
        "c3_muscle_ml": _round(_num(c3.get("class1_muscle_ml"))),
        "c3_brownfat_ml": _round(_num(c3.get("class2_brownfat_ml"))),
        "c3_mixwhite_ml": _round(_num(c3.get("class3_mixwhite_ml"))),
        "c3_total_ml": _round(_num(volumes.get("class3_total_ml"))),
        "c4_muscle_ml": _round(_num(c4.get("class1_muscle_ml"))),
        "c4_brownfat_ml": _round(_num(c4.get("class2_brownfat_ml"))),
        "c4_mixfat_ml": _round(_num(c4.get("class3_mixfat_ml"))),
        "c4_whitefat_ml": _round(_num(c4.get("class4_whitefat_ml"))),
        "c4_total_ml": _round(_num(volumes.get("class4_total_ml"))),

        "voxel_volume_ml": _round(_num(stats.get("voxel_volume_ml")), 6),
        "pred_voxels": stats.get("pred_voxels"),
    }

    return row


def describe(values: Iterable[Optional[float]]) -> dict:
    """n / mean / SD / min / max over the non-missing values."""
    vals = [v for v in values if v is not None]
    if not vals:
        return {"n": 0, "mean": None, "sd": None, "min": None, "max": None}
    return {
        "n": len(vals),
        "mean": round(mean(vals), 3),
        # Population SD: these are the cases in hand, not a sample drawn from a
        # larger population we are trying to infer about.
        "sd": round(pstdev(vals), 3) if len(vals) > 1 else 0.0,
        "min": round(min(vals), 3),
        "max": round(max(vals), 3),
    }


def summarise(rows: list[dict], total_cases: Optional[int] = None) -> dict:
    """Cohort-level aggregates: what a reviewer reads before the table.

    `total_cases` is every case in the system, which the rows cannot report --
    a row only carries measurements once its case has been analysed.
    """
    measured = [r for r in rows if r.get("has_metrics")]
    statuses: dict[str, int] = {}
    for row in rows:
        key = (row.get("status") or "unknown").upper()
        statuses[key] = statuses.get(key, 0) + 1

    return {
        "cases": len(rows),
        "total_cases": len(rows) if total_cases is None else total_cases,
        "measured": len(measured),
        "unmeasured": len(rows) - len(measured),
        "statuses": statuses,
        "metrics": {key: describe(r.get(key) for r in measured) for key in SUMMARY_KEYS},
    }


# --------------------------------------------------------------------------- #
#  export
# --------------------------------------------------------------------------- #
def rows_to_csv(rows: list[dict], columns=ALL_COLUMNS) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([header for _, header in columns])
    for row in rows:
        writer.writerow(["" if row.get(k) is None else row.get(k) for k, _ in columns])
    # utf-8-sig: Excel assumes the system codepage without a BOM and mangles
    # any non-ASCII patient name.
    return buffer.getvalue().encode("utf-8-sig")


def build_workbook(rows: list[dict], summary: dict, columns=ALL_COLUMNS) -> bytes:
    """Two sheets: the per-case table, and the cohort aggregates behind it."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cohort"

    sheet.append([header for _, header in columns])
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        sheet.append([row.get(key) for key, _ in columns])

    # Freeze the header and the subject column so scrolling right keeps the row
    # identifiable.
    sheet.freeze_panes = "B2"
    for index, (_, header) in enumerate(columns, start=1):
        width = max(11, min(len(header) + 3, 30))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = sheet.dimensions

    stats_sheet = workbook.create_sheet("Summary")
    stats_sheet.append(["Cases in selection", summary.get("cases")])
    stats_sheet.append(["With measurements", summary.get("measured")])
    stats_sheet.append(["Without measurements", summary.get("unmeasured")])
    stats_sheet.append([])
    stats_sheet.append(["Measure", "n", "Mean", "SD", "Min", "Max"])
    for cell in stats_sheet[5]:
        cell.font = header_font
        cell.fill = header_fill

    labels = dict(ALL_COLUMNS)
    for key in SUMMARY_KEYS:
        d = summary.get("metrics", {}).get(key, {})
        stats_sheet.append(
            [labels.get(key, key), d.get("n"), d.get("mean"), d.get("sd"),
             d.get("min"), d.get("max")]
        )
    stats_sheet.column_dimensions["A"].width = 26
    for column in "BCDEF":
        stats_sheet.column_dimensions[column].width = 12

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
